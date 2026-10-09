"""Exporta o orçamento para CSV (ponto e vírgula, padrão do Excel em português) e XLSX."""

import csv
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from orcamento.orcamento import LinhaOrcamento, Orcamento

CABECALHO = [
    "Situação",
    "Linha original",
    "Código SINAPI",
    "Descrição SINAPI",
    "Unidade",
    "Quantidade",
    "Preço unitário (R$)",
    "Subtotal (R$)",
    "Confiança",
    "Observação",
]


def _reais(centavos: int | None) -> Decimal | None:
    return None if centavos is None else Decimal(centavos) / 100


def _linha(item: LinhaOrcamento) -> list[object]:
    return [
        "REVISAR" if item.revisar else "OK",
        item.linha,
        item.codigo or "",
        item.descricao or "",
        item.unidade or item.unidade_pedida or "",
        item.quantidade_sinapi if item.quantidade_sinapi is not None else item.quantidade,
        _reais(item.preco_unitario_centavos),
        _reais(item.subtotal_centavos) if not item.revisar else None,
        item.confianca,
        item.motivo,
    ]


def _br(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, Decimal):
        return f"{valor:.2f}".replace(".", ",") if valor % 1 else f"{valor:.0f}"
    return str(valor)


def exportar_csv(orcamento: Orcamento, caminho: Path) -> None:
    with caminho.open("w", newline="", encoding="utf-8-sig") as arquivo:
        escritor = csv.writer(arquivo, delimiter=";")
        escritor.writerow(CABECALHO)
        for item in orcamento.itens:
            escritor.writerow([_br(v) for v in _linha(item)])
        escritor.writerow([])
        escritor.writerow(
            ["TOTAL (itens OK)", "", "", "", "", "", "", _br(_reais(orcamento.total_centavos))]
        )
        escritor.writerow(["Itens para revisar", len(orcamento.para_revisar)])


def exportar_xlsx(orcamento: Orcamento, caminho: Path) -> None:
    planilha = Workbook()
    aba = planilha.active
    assert aba is not None
    aba.title = "Orçamento"
    aba.append([f"Orçamento - SINAPI {orcamento.uf} {orcamento.referencia:%m/%Y}"])
    aba["A1"].font = Font(bold=True, size=13)
    aba.append([])
    aba.append(CABECALHO)
    for celula in aba[3]:
        celula.font = Font(bold=True)

    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for item in orcamento.itens:
        aba.append(_linha(item))
        if item.revisar:
            for celula in aba[aba.max_row]:
                celula.fill = amarelo

    aba.append([])
    aba.append(
        ["TOTAL (itens OK)", None, None, None, None, None, None, _reais(orcamento.total_centavos)]
    )
    aba[aba.max_row][0].font = Font(bold=True)
    aba[aba.max_row][7].font = Font(bold=True)
    aba.append(["Itens para revisar", len(orcamento.para_revisar)])

    for coluna in ("G", "H"):
        for celula in aba[coluna][3:]:
            celula.number_format = "#,##0.00"
    larguras = {
        "A": 10,
        "B": 34,
        "C": 13,
        "D": 60,
        "E": 9,
        "F": 11,
        "G": 18,
        "H": 15,
        "I": 11,
        "J": 60,
    }
    for coluna, largura in larguras.items():
        aba.column_dimensions[coluna].width = largura
    aba.freeze_panes = "A4"
    planilha.save(caminho)


def exportar(orcamento: Orcamento, caminho: Path) -> None:
    if caminho.suffix.lower() == ".xlsx":
        exportar_xlsx(orcamento, caminho)
    elif caminho.suffix.lower() == ".csv":
        exportar_csv(orcamento, caminho)
    else:
        raise ValueError("use .csv ou .xlsx")
